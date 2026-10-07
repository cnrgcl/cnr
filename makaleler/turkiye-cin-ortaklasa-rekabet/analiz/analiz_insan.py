# İnsan kodlayıcı analizi (K1, K2): Cohen κ, Ö1–Ö3 tabloları, ön-kayıtlı Ö3 kuralı, uyuşmazlık listesi.
# Girdi: analiz/insan-kodlar/K1.csv, K2.csv (ham xlsx'ten aktarım). Opsiyonel: analiz/insan-kodlar/uzlasi.csv
# (yazarın uyuşmazlık kararları; varsa "uzlaşı" seti nihai sonuç olur). Çıktı: analiz/stats-insan.json,
# analiz/insan-kodlar/uyusmazliklar.xlsx, analiz/sekil-verileri/*.csv (yazar şekilleri Excel'de çizer).
import csv, json, os, sys
from collections import Counter
sys.stdout.reconfigure(encoding="utf-8")

AR = ["1", "2", "3", "4"]
ADI = {"1": "Bağlantısallık", "2": "Turist akışı", "3": "Anlatı", "4": "Değer yakalama"}
TARAF = {"Türk kaynağı": "TR", "Çin kaynağı": "CN", "Ortak belge": "ORT"}

def oku(k):
    return {r["ifade_id"]: r for r in csv.DictReader(open(f"analiz/insan-kodlar/{k}.csv", encoding="utf-8"))}

def yonelim(r):
    m1, m2, m3 = r["m1"] == "E", r["m2"] == "E", r["m3"] == "E"
    if m1 and (m2 or m3): return "KR"
    if m1: return "İB"
    if m2 or m3: return "RK"
    return "NÖ"

def kappa(a, b):
    n = len(a); po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b); pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / n / n
    return {"n": n, "uyusma_yuzde": round(100 * po, 1), "kappa": round((po - pe) / (1 - pe), 3)}

def tablo(rows):
    out = {}
    for a in AR:
        c = Counter(yonelim(r) for r in rows if r["alan"] == a)
        n = sum(c.values()); nex = n - c["NÖ"]
        out[a] = {"n": n, "İB": c["İB"], "RK": c["RK"], "KR": c["KR"], "NÖ": c["NÖ"],
                  "rk_nö_haric": round(100 * c["RK"] / nex, 1) if nex else None,
                  "rk_nö_dahil": round(100 * c["RK"] / n, 1) if n else None,
                  "ib_nö_haric": round(100 * c["İB"] / nex, 1) if nex else None}
    return out

def o3(t, key):
    s = [t[a][key] for a in AR]
    if None in s: return {"seri": s, "karar": "hesaplanamadı"}
    mono = all(s[i] < s[i + 1] for i in range(3)); fark = round(s[3] - s[0], 1)
    return {"seri": s, "monoton": mono, "fark_A1_A4": fark,
            "karar": "destek" if (mono or fark >= 20) else ("destek yok" if fark <= 0 else "kısmi")}

def ozet(rows):
    t = tablo(rows); d = {"tablo": t, "o3_nö_haric": o3(t, "rk_nö_haric"), "o3_nö_dahil": o3(t, "rk_nö_dahil")}
    for tf in ("TR", "CN", "ORT"):
        tt = tablo([r for r in rows if TARAF[r["taraf"]] == tf]); d[f"tablo_{tf}"] = tt
        if tf != "ORT": d[f"o3_{tf}_nö_haric"] = o3(tt, "rk_nö_haric")
    for kp in ("A", "B"):
        sub = [r for r in rows if r["belge"].startswith(kp)]
        c = Counter(yonelim(r) for r in sub)
        d[f"korpus{kp}"] = {"n": len(sub), **dict(c), "ib_yuzde": round(100 * c["İB"] / len(sub), 1), "rk_yuzde": round(100 * c["RK"] / len(sub), 1)}
    for tf in ("TR", "CN"):
        sub = [r for r in rows if TARAF[r["taraf"]] == tf]
        d[f"isaret_{tf}"] = {"n": len(sub), "m2_kendini_one_cikarma": sum(r["m2"] == "E" for r in sub), "m3_dislama": sum(r["m3"] == "E" for r in sub)}
    return d

K1, K2 = oku("K1"), oku("K2")
ids = sorted(K1)
assert set(K1) == set(K2) and len(ids) == 368
st = {"n_birim": len(ids), "kodlayicilar": "K1, K2 (iki bağımsız insan kodlayıcı; kullanıcı beyanı 2026-10-07)"}
st["kappa"] = {
    "alan": kappa([K1[i]["alan"] for i in ids], [K2[i]["alan"] for i in ids]),
    "m1_ortak_ozne": kappa([K1[i]["m1"] for i in ids], [K2[i]["m1"] for i in ids]),
    "m2_kendini_one_cikarma": kappa([K1[i]["m2"] for i in ids], [K2[i]["m2"] for i in ids]),
    "m3_dislama": kappa([K1[i]["m3"] for i in ids], [K2[i]["m3"] for i in ids]),
    "yonelim": kappa([yonelim(K1[i]) for i in ids], [yonelim(K2[i]) for i in ids]),
}
st["yonelim_capraz"] = {f"{a}|{b}": c for (a, b), c in Counter((yonelim(K1[i]), yonelim(K2[i])) for i in ids).items()}
st["K1"] = ozet([K1[i] for i in ids]); st["K2"] = ozet([K2[i] for i in ids])
uyumlu = [K1[i] for i in ids if K1[i]["alan"] == K2[i]["alan"] and yonelim(K1[i]) == yonelim(K2[i])]
st["uyumlu_alt_kume"] = {"n": len(uyumlu), **ozet(uyumlu)}

# Yazar uzlaşısı varsa nihai set
if os.path.exists("analiz/insan-kodlar/uzlasi.csv"):
    uz = {r["ifade_id"]: r for r in csv.DictReader(open("analiz/insan-kodlar/uzlasi.csv", encoding="utf-8"))}
    nihai = []
    for i in ids:
        r = dict(K1[i])
        if i in uz and uz[i].get("karar_alan"):
            r.update(alan=uz[i]["karar_alan"], m1=uz[i]["karar_m1"], m2=uz[i]["karar_m2"], m3=uz[i]["karar_m3"])
        elif not (K1[i]["alan"] == K2[i]["alan"] and yonelim(K1[i]) == yonelim(K2[i])):
            r = None
        if r: nihai.append(r)
    st["nihai_uzlasi"] = {"n": len(nihai), **ozet(nihai)}

json.dump(st, open("analiz/stats-insan.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)

# Uyuşmazlık listesi (yazar karar verir)
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
metin = {r["ifade_id"]: r["metin"] for r in csv.DictReader(open("analiz/coklu-kodlayici-ifade-esleme.csv", encoding="utf-8"))}
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Uyuşmazlıklar"
bas = ["İfade no", "Taraf", "İfade", "K1 alan", "K1 M1", "K1 M2", "K1 M3", "K1 yönelim", "K2 alan", "K2 M1", "K2 M2", "K2 M3", "K2 yönelim",
       "KARAR alan", "KARAR M1", "KARAR M2", "KARAR M3", "Gerekçe"]
ws.append(bas)
for c in ws[1]: c.font = Font(bold=True)
sari = PatternFill("solid", fgColor="FFF2CC")
for i in ids:
    a, b = K1[i], K2[i]
    if a["alan"] != b["alan"] or yonelim(a) != yonelim(b):
        ws.append([i, a["taraf"], metin.get(i, ""), a["alan"], a["m1"], a["m2"], a["m3"], yonelim(a), b["alan"], b["m1"], b["m2"], b["m3"], yonelim(b), "", "", "", "", ""])
        for col in range(14, 19): ws.cell(ws.max_row, col).fill = sari
ws.column_dimensions["C"].width = 70
for row in ws.iter_rows(min_row=2):
    row[2].alignment = Alignment(wrap_text=True, vertical="top")
wb.save("analiz/insan-kodlar/uyusmazliklar.xlsx")

# Şekil verileri (yazar Excel'de çizer)
os.makedirs("analiz/sekil-verileri", exist_ok=True)
with open("analiz/sekil-verileri/sekil3-rekabet-payi.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(["alan", "K1_rk_nö_haric", "K2_rk_nö_haric", "uyumlu_rk_nö_haric", "uyumlu_TR", "uyumlu_CN"])
    for a in AR:
        w.writerow([ADI[a], st["K1"]["tablo"][a]["rk_nö_haric"], st["K2"]["tablo"][a]["rk_nö_haric"],
                    st["uyumlu_alt_kume"]["tablo"][a]["rk_nö_haric"], st["uyumlu_alt_kume"]["tablo_TR"][a]["rk_nö_haric"],
                    st["uyumlu_alt_kume"]["tablo_CN"][a]["rk_nö_haric"]])
print(json.dumps({"kappa": st["kappa"], "uyumlu_n": st["uyumlu_alt_kume"]["n"],
                  "o3": {k: st[k]["o3_nö_haric"] for k in ("K1", "K2", "uyumlu_alt_kume")},
                  "o3_dahil": {k: st[k]["o3_nö_dahil"] for k in ("K1", "K2", "uyumlu_alt_kume")}}, ensure_ascii=False, indent=1))
